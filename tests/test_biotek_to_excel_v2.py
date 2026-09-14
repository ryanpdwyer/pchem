"""Tests for biotek_to_excel_v2.py (combine + tidy reshape).

Fixtures are generated programmatically so we can place a run in any plate
column, which is what the combine step keys on. Run with:
    micromamba run -n py314 python -m pytest -q test_biotek_to_excel_v2.py
"""
from __future__ import annotations

import os
import tempfile

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from pchemapps import biotek_to_excel as v1
from pchemapps import biotek_to_excel_v2 as v2


# --- Fixture generator -------------------------------------------------------


def _grid_line(letter: str, col: int, value: float) -> str:
    cells = [""] * 13
    cells[0] = letter
    cells[col] = f"{value:.3f}"
    return "\t".join(cells)


def _kinetic_line(hms: str, col_idx0: int, value: float) -> str:
    """col_idx0 is the 0-based well index into WELLS for row A of `col`."""
    vals = [""] * 96
    vals[col_idx0] = f"{value:.3f}"
    return hms + "\t" + "\t".join(vals)


def make_run(col: int, init734=0.500, init950=0.080,
             k734=(0.400, 0.300, 0.200), k950=(0.050, 0.052, 0.054),
             times=("0:00:00", "0:01:19", "0:02:38")) -> str:
    """A one-well (row A), one-column full kinetic run in Gen5 text form."""
    a_idx = v1.WELLS.index(f"A{col}")  # row-A well index for this column
    lines = [
        "Procedure Summary", "",
        "Plate Type\tCostar 96 flat bottom",
        "Read\tRead:  Intial Read before antioxidant (A) 734, 950",
        "Start Kinetic\tStart Kinetic [Run 0:06:20, Interval 0:01:19]", "",
        "Intial Read before antioxidant:734",
        "\t1\t2\t3\t4\t5\t6\t7\t8\t9\t10\t11\t12",
        _grid_line("A", col, init734), "",
        "Intial Read before antioxidant:950",
        "\t1\t2\t3\t4\t5\t6\t7\t8\t9\t10\t11\t12",
        _grid_line("A", col, init950), "",
        "734,950:734", "",
        "Kinetic read\t" + "\t".join(v1.WELLS),
    ]
    lines += [_kinetic_line(t, a_idx, v) for t, v in zip(times, k734)]
    lines += ["", "734,950:950", "", "Kinetic read\t" + "\t".join(v1.WELLS)]
    lines += [_kinetic_line(t, a_idx, v) for t, v in zip(times, k950)]
    return "\n".join(lines) + "\n"


def _write_temp(content: str) -> str:
    fd, path = tempfile.mkstemp(suffix=".txt")
    with os.fdopen(fd, "wb") as f:
        f.write(content.encode("cp1252"))
    return path


def _plate_info(labels: dict[int, str]) -> pd.DataFrame:
    """Blank PlateInfo with row-A Sample wells labeled per {column: sample}."""
    df = v1.make_blank_plate_info()
    for col, sample in labels.items():
        loc = f"A{col}"
        df.loc[df["Loc"] == loc, ["Type", "Sample", "Conc_mM"]] = ["Sample", sample, 0.5]
    return df


# --- derive_serial_dilution --------------------------------------------------


def test_derive_serial_dilution_halves_down_column():
    df = v1.make_blank_plate_info()
    # Column 1: A=Sample 0.5, B-F Sample blank, G Control 0, H Blank.
    for r, typ in zip("ABCDEF", ["Sample"] * 6):
        df.loc[df["Loc"] == f"{r}1", "Type"] = typ
    df.loc[df["Loc"] == "A1", "Conc_mM"] = 0.5
    df.loc[df["Loc"] == "G1", ["Type", "Conc_mM"]] = ["Control", 0]
    df.loc[df["Loc"] == "H1", "Type"] = "Blank (EtOH)"

    out = v2.derive_serial_dilution(df)
    g = out.set_index("Loc")
    assert g.loc["A1", "Conc_mM"] == pytest.approx(0.5)
    assert g.loc["B1", "Conc_mM"] == pytest.approx(0.25)
    assert g.loc["C1", "Conc_mM"] == pytest.approx(0.125)
    assert g.loc["F1", "Conc_mM"] == pytest.approx(0.5 / 32)
    # Anchor not flagged derived; filled ones are.
    assert not g.loc["A1", "Conc_derived"]
    assert g.loc["B1", "Conc_derived"]
    # Control left untouched (still 0, not derived).
    assert g.loc["G1", "Conc_mM"] == pytest.approx(0.0)
    assert not g.loc["G1", "Conc_derived"]


# --- detect_data_column ------------------------------------------------------


def test_detect_data_column_finds_the_used_column():
    for col in (1, 4, 7):
        path = _write_temp(make_run(col))
        try:
            pf = v1.parse_file(path)
            assert v2.detect_data_column(pf) == [col]
        finally:
            os.unlink(path)


# --- build_tidy --------------------------------------------------------------


def test_build_tidy_has_wavelength_columns_and_flags():
    path = _write_temp(make_run(1))
    try:
        pf = v1.parse_file(path)
        pi = v2.derive_serial_dilution(_plate_info({1: "Trolox"}))
        tidy = v2.build_tidy(pf, pi, "Trolox")
        assert set(["A_734", "A_950", "dA", "is_initial", "is_final"]) <= set(tidy.columns)

        # Initial read row: A_734=0.500, A_950=0.080 -> dA=0.420, is_initial True.
        init = tidy[(tidy["Loc"] == "A1") & (tidy["Time_index"] == -1)].iloc[0]
        assert init["A_734"] == pytest.approx(0.500)
        assert init["dA"] == pytest.approx(0.420)
        assert bool(init["is_initial"]) and not bool(init["is_final"])

        # Last kinetic point (index 2) is the final, not initial.
        fin = tidy[(tidy["Loc"] == "A1") & (tidy["Time_index"] == 2)].iloc[0]
        assert fin["dA"] == pytest.approx(0.200 - 0.054)
        assert bool(fin["is_final"]) and not bool(fin["is_initial"])
        # Exactly one final and one initial per well.
        a1 = tidy[tidy["Loc"] == "A1"]
        assert a1["is_final"].sum() == 1
        assert a1["is_initial"].sum() == 1
    finally:
        os.unlink(path)


# --- combine + workbook ------------------------------------------------------


def test_combine_two_experiments_and_skips_unmatched(tmp_path):
    f1 = _write_temp(make_run(1, init734=0.50, k734=(0.40, 0.20)))
    f2 = _write_temp(make_run(2, init734=0.45, k734=(0.42, 0.30)))
    f_un = _write_temp(make_run(9))  # column 9 not labeled -> skipped
    try:
        pi = v2.derive_serial_dilution(_plate_info({1: "Trolox", 2: "Resorcinal"}))
        tidy, meta = v2.combine([f1, f2, f_un], pi)
        assert set(tidy["Experiment"].unique()) == {
            os.path.splitext(os.path.basename(f1))[0],
            os.path.splitext(os.path.basename(f2))[0],
        }
        # Summary records the skip.
        joined = " ".join(f"{k} {v}" for k, v in meta)
        assert "skipped" in joined

        out = tmp_path / "Combined.xlsx"
        v2.write_combined_workbook(tidy, pi, meta, str(out))
        wb = load_workbook(out)
        assert {"Summary", "PlateInfo", "Layout", "Data_long",
                "Per_well", "QC_initial", "dA_vs_time"} <= set(wb.sheetnames)
    finally:
        for p in (f1, f2, f_un):
            os.unlink(p)


def test_per_well_delta_is_initial_minus_final(tmp_path):
    f1 = _write_temp(make_run(1, init734=0.50, init950=0.08,
                              k734=(0.40, 0.20), k950=(0.05, 0.06)))
    try:
        pi = v2.derive_serial_dilution(_plate_info({1: "Trolox"}))
        tidy, _ = v2.combine([f1], pi)
        pw = v2.build_per_well(tidy)
        row = pw[pw["Loc"] == "A1"].iloc[0]
        # dA_initial = 0.50-0.08 = 0.42; dA_final = 0.20-0.06 = 0.14.
        assert row["dA_initial"] == pytest.approx(0.42)
        assert row["dA_final"] == pytest.approx(0.14)
        assert row["delta_dA"] == pytest.approx(0.28)
        assert row["frac_remaining"] == pytest.approx(0.14 / 0.42)
    finally:
        os.unlink(f1)


def test_dilution_factor_adds_in_well_concentration(tmp_path):
    """Conc_mM is what was pipetted; Conc_final_mM is what the well saw.

    Demmi's assay adds 10 uL of antioxidant to 190 uL of radical solution, so
    the plate map's 10 mM is 0.5 mM in the well -- a c_50 read off Conc_mM
    would be 20x too high.
    """
    f1 = _write_temp(make_run(1, init734=0.50, init950=0.08,
                              k734=(0.40, 0.20), k950=(0.05, 0.06)))
    try:
        pi = v2.derive_serial_dilution(_plate_info({1: "Trolox"}))

        tidy, meta = v2.combine([f1], pi, dilution=20.0)
        assert "Conc_final_mM" in tidy.columns
        row = tidy[tidy["Loc"] == "A1"].iloc[0]
        assert row["Conc_mM"] == pytest.approx(0.5)
        assert row["Conc_final_mM"] == pytest.approx(0.025)
        assert v2.build_per_well(tidy).iloc[0]["Conc_final_mM"] == pytest.approx(0.025)
        assert any("Conc_final_mM" in str(k) for k, _ in meta)

        # Default is a no-op: the column stays out of everyone else's workbooks.
        plain, plain_meta = v2.combine([f1], pi)
        assert "Conc_final_mM" not in plain.columns
        assert "Conc_final_mM" not in v2.build_per_well(plain).columns
        assert not any("Conc_final_mM" in str(k) for k, _ in plain_meta)
    finally:
        os.unlink(f1)
